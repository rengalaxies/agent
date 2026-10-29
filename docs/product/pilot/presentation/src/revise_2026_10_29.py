"""Revision 29.10.2026 of the eight-slide deck: facts of the delivered service, not the concept.

Applies text edits to the 24.10 PPTX built by build.mjs, keeping its layout and styles,
writes the notes from slides.json into the PPTX and regenerates SPEAKER-NOTES.md.
Needs python-pptx. PDF and slide previews come from the same PPTX (export_pdf.ps1).

Usage: python revise_2026_10_29.py <source.pptx> <output.pptx>
"""
import copy
import json
from pathlib import Path
import sys

from pptx import Presentation

# slide number -> {shape name prefix: lines}; one line per paragraph.
EDITS = {
    1: {"Концепция автономного MVP": ["Закрытый технический пилот · 29 октября 2026"]},
    4: {"Трассы и журналы": ["Журналы доступа"]},
    6: {"Готовый образ": ["Релиз и systemd", "на Linux-ВМ"],
        "Код, вызовы": ["GitHub App, Jira", "HTTP/Python-код"],
        "Модель, доступы": ["Доступы и лимиты", "модель выключена"],
        "В штатном потоке оператор": ["В штатном потоке оператор не нужен.",
                                      "Сбои, свежесть и пределы анализа видны в комментарии."]},
    7: {"Slide title": ["Что проверено и что ещё нет"],
        "Работает в прототипе": ["Проверено локально и в CI"],
        "Узкий статический анализ": ["Jira и PR → один комментарий",
                                     "Основание: репо@коммит:строка",
                                     "Установка и откат релиза (CI systemd)",
                                     "Показ: изменение → исправление"],
        "Следующий автономный выпу": ["Ещё не проверено"],
        "Jira и PR без ручного зап": ["Настоящие GitHub App и Jira (LIVE)",
                                      "Установка на ВМ Cloud.ru",
                                      "Незнакомый проект и польза команды",
                                      "Модель: допуск не пройден, OFF"],
        "Преимущество, экономия": ["Профиль пилота: HTTP API на Python. Экономия и преимущество не измерены."]},
    8: {"Slide title": ["Закрытый пилот: один сквозной сценарий"],
        "Один API и его потребител": ["Один HTTP/Python API и потребители",
                                      "2–5 репозиториев GitHub",
                                      "Проект Jira, если он есть",
                                      "Права: чтение и свой комментарий"],
        "Время до полезного инсайт": ["Время до полезного вывода",
                                      "Пропуски и ложные сигналы",
                                      "Ручные минуты подключения",
                                      "Сравнение с обычным процессом"],
        "Цель — полезный ответ ран": ["Условия пилота: предложение с допущениями,",
                                      "а не утверждённая сделка."]},
}


def set_lines(shape, lines):
    """Replace paragraph texts in place; the first run of each paragraph keeps its style."""
    body = shape.text_frame._txBody
    paragraphs = list(shape.text_frame.paragraphs)
    while len(paragraphs) < len(lines):
        body.append(copy.deepcopy(paragraphs[-1]._p))
        paragraphs = list(shape.text_frame.paragraphs)
    for paragraph, line in zip(paragraphs, lines):
        runs = paragraph.runs
        runs[0].text = line
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
    for paragraph in paragraphs[len(lines):]:
        body.remove(paragraph._p)


def speaker_notes(slides) -> str:
    """SPEAKER-NOTES.md from slides.json, the same text as the PPTX notes."""
    out = ["# Заметки к восьми слайдам", "",
           "Редакция 29.10.2026: факты поставленного сервиса и условия закрытого пилота.", ""]
    for s in slides:
        out += [f"## {s['slide']}. {s['title'] or 'Кого затронет следующее изменение?'}", "", s["notes"], "",
                "Источники:"] + [f"- {src}" for src in s["sources"]] + [""]
    return "\n".join(out)


def revise(source, output, slides_json):
    slides = json.loads(Path(slides_json).read_text(encoding="utf-8"))
    deck = Presentation(source)
    for number, edits in EDITS.items():
        slide = deck.slides[number - 1]
        for prefix, lines in edits.items():
            [shape] = [s for s in slide.shapes if s.has_text_frame and s.name.startswith(prefix)]
            set_lines(shape, lines)
    for item, slide in zip(slides, deck.slides):
        slide.notes_slide.notes_text_frame.text = item["notes"]
    deck.save(output)
    notes = Path(slides_json).parents[1] / "SPEAKER-NOTES.md"
    notes.write_text(speaker_notes(slides), encoding="utf-8", newline="\n")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    revise(sys.argv[1], sys.argv[2], Path(__file__).with_name("slides.json"))
